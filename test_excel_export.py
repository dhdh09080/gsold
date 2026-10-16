import io
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd
import streamlit as st
from streamlit.testing.v1 import AppTest

from excel_export import create_excel, measurement_date_label, measurement_title


class ExcelExportTests(unittest.TestCase):
    def test_print_layout_and_content(self):
        start, end = date(2026, 10, 19), date(2026, 10, 20)
        roster = pd.DataFrame({
            'No.': [1], '협력회사': ['테스트회사'], '이름': ['테스트'], '공종': ['목공'],
            '혈압잰 일자': [measurement_date_label(start, end)],
            '오늘의 혈압': ['     /     '], '고혈압 복용 여부': ['□예  □아니오'], '서명란': [''],
        })
        sheet = openpyxl.load_workbook(create_excel(roster, start, end)).active
        self.assertEqual(sheet['A1'].value, '2026.10.19~10.20 고령자 혈압측정')
        self.assertEqual(str(sheet.merged_cells), 'A1:H1')
        self.assertEqual([c.value for c in sheet[2]], list(roster.columns))
        self.assertEqual(sheet['E3'].value, '2026년 10월      일')
        self.assertEqual(sheet['F3'].value, '     /     ')
        self.assertEqual(sheet.print_title_rows, '$1:$2')
        self.assertEqual(str(sheet.page_setup.paperSize), sheet.PAPERSIZE_A4)
        # XLSX omits fitToWidth when it uses Excel's default of one page.
        self.assertIn(sheet.page_setup.fitToWidth, (None, 1))
        self.assertEqual(sheet.page_setup.fitToHeight, 0)
        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(sheet['A1'].font.sz, 18)
        for row in sheet.iter_rows(min_row=2, max_row=13, max_col=8):
            for cell in row:
                self.assertEqual(cell.font.sz, 12)
        for row in range(1, 14):
            self.assertEqual(sheet.row_dimensions[row].height, 25)
        self.assertEqual(sheet['H13'].border.bottom.style, 'thin')
        self.assertEqual(sheet['G3'].value, '□예  □아니오')
        self.assertEqual(sheet['G13'].value, '□예  □아니오')

    def test_single_day_and_month_year_boundaries(self):
        day = date(2026, 10, 19)
        self.assertEqual(measurement_title(day, day), '2026.10.19 고령자 혈압측정')
        self.assertEqual(measurement_date_label(day, day), '2026년 10월 19일')
        self.assertEqual(measurement_title(date(2026, 10, 31), date(2026, 11, 1)),
                         '2026.10.31~11.01 고령자 혈압측정')
        self.assertEqual(measurement_title(date(2026, 12, 31), date(2027, 1, 1)),
                         '2026.12.31~2027.01.01 고령자 혈압측정')
        self.assertEqual(measurement_date_label(date(2026, 10, 31), date(2026, 11, 1)),
                         '        년     월     일')
        with self.assertRaises(ValueError):
            measurement_title(date(2026, 10, 20), day)

    def test_app_date_selection_changes_download(self):
        app = AppTest.from_file(str(Path(__file__).with_name('app.py'))).run(timeout=30)
        app.date_input[0].set_value(date(2026, 10, 19))
        app.date_input[1].set_value(date(2026, 10, 20))
        rows = []
        for name, birthday in [('고령근로자', '500101-1'), ('청년근로자', '000101-3')]:
            row = [''] * 12
            row[3], row[4], row[5], row[9], row[11] = '테스트회사', '목공', name, '한국', birthday
            rows.append(row)
        uploaded = io.BytesIO()
        pd.DataFrame(rows).to_excel(uploaded, index=False)
        uploaded.seek(0)
        downloads = {}
        original_download = st.download_button

        def capture_download(**kwargs):
            downloads[kwargs['file_name']] = kwargs['data'].getvalue()
            return original_download(**kwargs)

        with patch('streamlit.file_uploader', return_value=uploaded), \
                patch('streamlit.download_button', side_effect=capture_download):
            app.run(timeout=60)
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        self.assertEqual([metric.value for metric in app.metric], ['2 명', '1 명', '0 명'])
        self.assertEqual(app.dataframe[0].value['이름'].tolist(), ['고령근로자'])
        name = '2026.10.19~10.20 고령자 혈압측정'
        sheet = openpyxl.load_workbook(io.BytesIO(downloads[name + '.xlsx'])).active
        self.assertEqual(sheet['A1'].value, name)
        self.assertEqual(sheet['F3'].value, '     /     ')
        self.assertTrue(downloads[name + '.png'].startswith(b'\x89PNG\r\n\x1a\n'))
        app.date_input[1].set_value(date(2026, 10, 18)).run()
        self.assertFalse(app.exception)
        self.assertIn('종료일', app.warning[0].value)
        self.assertEqual(len(app.get('download_button')), 0)


if __name__ == '__main__':
    unittest.main()
